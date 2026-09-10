from __future__ import annotations

import csv
import hashlib
import re
import secrets
import string
import unicodedata
import uuid
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(r"C:\Users\pc\Downloads\Grupos y administradores por regiones.xlsx")
AGENCIES_SQL = ROOT / "database-mssql" / "005_import_agencies_ltk.sql"
OUTPUT_SQL = ROOT / "database-mssql" / "008_import_group_admins.sql"
CREDENTIALS_CSV = ROOT / "deploy" / "CREDENCIALES_ADMINISTRADORES_GRUPOS.csv"
ASSIGNMENTS_CSV = ROOT / "deploy" / "ASIGNACIONES_ADMINISTRADORES_GRUPOS.csv"


def clean(value: object | None) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized(value: str) -> str:
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", ".", normalized(value)).strip(".")[:42] or "usuario"


def sql_text(value: str | None) -> str:
    if not value:
        return "NULL"
    return "N'" + value.replace("'", "''") + "'"


def password() -> str:
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789!@#"
    while True:
        value = "".join(secrets.choice(alphabet) for _ in range(14))
        if any(c.isupper() for c in value) and any(c.islower() for c in value) and any(c.isdigit() for c in value):
            return value


source_lines = AGENCIES_SQL.read_text(encoding="utf-8").splitlines()
group_counts: Counter[str] = Counter()
for line in source_lines:
    if not line.startswith("(N'LTK:"):
        continue
    values = re.findall(r"N'((?:''|[^'])*)'", line)
    if len(values) == 4:
        group_counts[values[3].replace("''", "'")] += 1

actual_groups = sorted(group_counts)
if len(actual_groups) != 169 or sum(group_counts.values()) != 3319:
    raise RuntimeError(f"Catálogo inesperado: {len(actual_groups)} grupos y {sum(group_counts.values())} agencias")

book = load_workbook(WORKBOOK, read_only=True, data_only=True)
sheet = book["CONTACTOS"]
current_region = ""
barahona_section = False
rows: list[dict[str, str]] = []

for number, group_label, administrator, contact in sheet.iter_rows(min_row=1, max_col=4, values_only=True):
    label = clean(group_label)
    admin = clean(administrator)
    phone = clean(contact)
    heading = clean(number)
    if heading and not label and not admin:
        if "Región Este" in heading:
            current_region, barahona_section = "Región Este", False
        elif "Norte y Noreste" in heading:
            current_region, barahona_section = "Región Norte y Noreste", False
        elif heading == "Región Sur":
            current_region, barahona_section = "Región Sur", False
        elif "Grupos Barahona" in heading:
            current_region, barahona_section = "Región Sur · Barahona", True
        elif heading == "Gran Santo Domingo":
            current_region, barahona_section = "Gran Santo Domingo", False
        elif "Grupos Administrativos" in heading:
            current_region, barahona_section = "Grupos Administrativos", False
        continue
    if not isinstance(number, (int, float)) or not label:
        continue
    if barahona_section and not admin:
        admin = "Luis Manuel"
    if not admin or admin == "---":
        continue
    display_name = re.sub(r"\s*\((?:Encargada|Asistente)\)\s*", "", admin, flags=re.I).strip()
    phone_digits = re.sub(r"\D", "", phone)
    identity = f"phone:{phone_digits}" if len(phone_digits) >= 10 else f"name:{normalized(display_name)}"
    rows.append({"identity": identity, "name": display_name, "contact": phone if phone != "---" else "", "region": current_region, "label": label})

admins: dict[str, dict] = {}
label_owner: dict[str, str] = {}
for row in rows:
    record = admins.setdefault(row["identity"], {"name": row["name"], "contact": row["contact"], "regions": set(), "labels": set(), "groups": set()})
    record["regions"].add(row["region"])
    record["labels"].add(row["label"])
    if not record["contact"] and row["contact"]:
        record["contact"] = row["contact"]
    label_owner[normalized(row["label"])] = row["identity"]


def groups_for(*names: str) -> list[str]:
    missing = [name for name in names if name not in group_counts]
    if missing:
        raise RuntimeError("Grupos inexistentes en la asignación: " + ", ".join(missing))
    return list(names)


mapping: dict[str, list[str]] = {
    "Mario Romana": groups_for("0207 - Mario Romana"),
    "Kiko Romana": groups_for("209-KIKO ROMANA"),
    "Romana": groups_for("027 -ROMANA", "LOTTO REAL ROMANA", "ROMANA | FELO", "LUIS | RIFERO ROMANA"),
    "Higuey": groups_for("2-AMARILYS | HIGUEY", "HIGUEY DEPORTIVAS"),
    "Grupos Bavaro / Lumismart Punta Cana": groups_for("130-LUMISMART BAVARO | PUNTA CANA", "BAVARO B", "LOTTO REAL PUNTA CANA"),
    "Nisibon": groups_for("126-NISIBON N"),
    "Miches": groups_for("179-MICHES M"),
    "Yuma": groups_for("147-YUMA"),
    "Bayahibe": groups_for("132-LUMISMART BAYAHIBE BENERITO"),
    "Sanchez": groups_for("SANCHEZ", "LOTTO REAL SANCHEZ"),
    "Mao": groups_for("164-MAO VALVERDE"),
    "Monte Plata": groups_for("17-MONTE PLATA"),
    "Bayaguana": groups_for("143-BAYAGUANA SOCIEDAD FM RICHARD", "BAYAGUANA SPORT"),
    "Yamasá": groups_for("144-LUMISMART KM YAMASA", "SOCIEDAD SANDRA YAMASA"),
    "Samaná": groups_for("SAMANA", "LOTTO REAL SAMANA", "SOCIEDAD COLOMBO/ SAMANA"),
    "Hedwing Samaná": groups_for("HEDWING | SAMANA RIFERO"),
    "Kinito Samana": groups_for("196-KINITO | SAMANA RIFERO"),
    "Bani Roberto": groups_for("138-BANI"),
    "Bujia Barahona": groups_for("66-BUJIA | BARAHONA"),
    "Alberto Pedernales y Nelson Pedernales": groups_for("ALBERTO PEDERNALES", "18-NELSON PEDERNALES"),
    "Jimani Cocha": groups_for("TIERRA NUEVA JIMANI COCHA"),
    "Las Matas": groups_for("29-LAS MATAS FLOR MARIA"),
    "Hondo Valle": groups_for("54-HONDO VALLE MACHO FM", "HONDO VALLE MACHO ARRENDADAS"),
    "El Cercado Benito": groups_for("57-EL CERCADO BENITO"),
    "Azua": groups_for("171-GRUPO AZUA"),
    "Barahona Cabral Pachi": groups_for("69-BARAHONA CABRAL | PACHI", "105-BARAHONA | CABRAL"),
    "Barahona Vicente Nobles": groups_for("77-BARAHONA VICENTE NOBLES | TAMAYO"),
    "Aeropuertos / Metro M / Sociedad Mercadal/Carrefour / Sambil/ Joselito HS": groups_for("AEROPUERTO | AILA | HIGUERO", "METRO M", "39-SOCIEDAD MERCADAL GLEDYS, NANCY, ALBERTO", "112-CARREFOUR", "16-JOSELITO | HS | ALBERTO"),
    "Alberto propias": groups_for("3-ALBERTO PROPIAS", "ALBERTO SPORT"),
    "Lumismart Capital/ Grupo Eduard Guzman / Papo Estrella": groups_for("145-LUMISMART CAPITAL", "204-GRUPO EDUARD GUZMAN", "205-PAPO ESTRELLA"),
    "Carlos Brisas": groups_for("63-CARLOS BRIZAS"),
    "Compadre Jhony / Soñadora Hs / Jhon Silvestre": groups_for("43-COMPADRE JHONNY C", "SONADORA HS", "203-JHONATAN SILVESTRE HIJO / SOCIEDAD"),
    "David asociados": groups_for("DAVID | ASOCIADO"),
    "David Propias": groups_for("74-DAVID | PROPIAS"),
    "Divanna Boca chica": groups_for("4-DIVANNA BOCA CHICA"),
    "Dominga propia / Luis Alberto Capotillo": groups_for("96-DOMINGA PROPIA", "5-LUIS ALBERTO/ CAPOTILLO PROPIA"),
    "Domingo Mameyes": groups_for("41-DOMINGO MAMEYES"),
    "Dr. Enerio": groups_for("85-DR.ENERIO E"),
    "Eddy Junior": groups_for("72-EDDY JUNIOR"),
    "El príncipe": groups_for("45-EL PRINCIPE P"),
    "Fernando": groups_for("58-FERNANDO"),
    "Fortuna F": groups_for("127-FORTUNA F"),
    "Franklin Mameyes": groups_for("35-FRANKLIN MAMEYES"),
    "Franklin Reyes / Riferos FM": groups_for("7-FRANKLIN REYES | SOCIEDAD TEJEDA", "59-RIFEROS_FM_SOCIEDAD"),
    "Gallo G / Gallo 2": groups_for("92-GALLO G"),
    "Grupo JAT / Yordy": groups_for("201-GRUPO - J.A.T", "JHORDY SPORT"),
    "Grupo Santos Martinez": groups_for("GRUPO SANTO"),
    "Javier tejeda": groups_for("137-JAVIER TEJEDA"),
    "J&L": groups_for("195-JYL SOCIEDAD"),
    "Joel propias / Ozama": groups_for("JOEL | PROPIAS", "019 - OZAMA"),
    "Jose mesa": groups_for("53-JOSE MESA J"),
    "Jose Ramon sabana p": groups_for("26-JOSE RAMON | SABANA P."),
    "Valiente": groups_for("165-VALIENTE SDE"),
    "Los mina pedrito": groups_for("LOS MINA PEDRITO"),
    "Luis Tejeda DN": groups_for("LUIS TEJEDA .D.N"),
    "Mango / Richard Jumbo": groups_for("13-RICHARD MANGO M", "28-RICHARD JUMBO"),
    "Mechi invivienda": groups_for("48-MECHI INVIVIENDA"),
    "Miguel Rodriguez": groups_for("210-MIGUEL RODRIGUEZ ( LA FAMOSA)"),
    "Nelson bretón": groups_for("50-NELSON BRETON"),
    "Ney": groups_for("NEY"),
    "Peje capital / Peje Castillos": groups_for("62-PEJE | CAPITAL", "PEJE | CASTILLO"),
    "Arismedy (La Fortuna)": groups_for("ARISMENDI LA FORTUNA"),
    "Simon S": groups_for("015 - Grupo Simon"),
    "Tejeda Pimentel SDE": groups_for("197-TEJEDA PIMENTEL| SDE,"),
    "Tio Máximo": groups_for("89-EL TIO MAXIMO"),
    "Villa mella": groups_for("VILLA MELLA V"),
    "Wascar SP": groups_for("146-WASCAR SP"),
    "Xiomara": groups_for("61-XIOMARA X"),
    "Sociedad Dario Alberto D": groups_for("135-SOCIEDAD DARIO | ALBERTO D"),
    "Bryan B / Bryan Luis Alberto": groups_for("BRYAN B", "BRYAN | LUIS ALBERTO"),
    "Sociedad Arias": groups_for("118-SOCIEDAD ARIAS TEJEDA"),
}

# Los encabezados amplios del Excel representan varios grupos reales.
mapping["Grupos Santiago"] = [g for g in actual_groups if "SANTIAGO" in g or "SAJOMA" in g]
mapping["Constanza"] = [g for g in actual_groups if "SURIEL" in g or "CONSTANZA" in g or "JARABACOA" in g]

# Barahona hereda a Luis Manuel salvo los grupos con responsable explícito.
luis_identity = label_owner.get(normalized("Barahona Centro de Ciudad"))
if not luis_identity:
    raise RuntimeError("No se pudo identificar a Luis Manuel en la sección Barahona")
for group in actual_groups:
    if "BARAHONA" in group or group in {"NEIBA 1", "AGENCIAS SPORT"}:
        admins[luis_identity]["groups"].add(group)

group_owner: dict[str, str] = {group: luis_identity for group in admins[luis_identity]["groups"]}
for label, groups in mapping.items():
    owner = label_owner.get(normalized(label))
    if not owner:
        raise RuntimeError(f"El Excel no contiene un administrador utilizable para: {label}")
    for group in groups:
        previous = group_owner.get(group)
        if previous and previous != owner:
            admins[previous]["groups"].discard(group)
        group_owner[group] = owner
        admins[owner]["groups"].add(group)

# Cuentas únicas. Cuando hay homónimos, el teléfono evita colisiones.
used_emails: set[str] = set()
for identity, admin in sorted(admins.items(), key=lambda item: normalized(item[1]["name"])):
    base = slug(admin["name"])
    email = f"{base}@grupotejeda.local"
    if email in used_emails:
        digits = re.sub(r"\D", "", admin["contact"])
        suffix = digits[-4:] if digits else str(len(used_emails) + 1)
        email = f"{base}.{suffix}@grupotejeda.local"
    used_emails.add(email)
    admin["email"] = email
    admin["id"] = str(uuid.uuid4())
    admin["password"] = password()
    admin["salt"] = secrets.token_bytes(16)
    admin["hash"] = hashlib.pbkdf2_hmac("sha256", admin["password"].encode("utf-8"), admin["salt"], 210000, 32)
    admin["region"] = " · ".join(sorted(admin["regions"]))

sql: list[str] = [
    "-- Administradores por regiones y permisos por grupo - SQL Server 2025",
    "-- Fuente: Grupos y administradores por regiones.xlsx",
    "-- Las contraseñas temporales están fuera de wwwroot y deben cambiarse al primer ingreso.",
    "",
    "SET NOCOUNT ON;",
    "SET XACT_ABORT ON;",
    "",
    "IF COL_LENGTH(N'dbo.admin_users',N'region') IS NULL ALTER TABLE dbo.admin_users ADD region nvarchar(120) NULL;",
    "IF COL_LENGTH(N'dbo.admin_users',N'contact') IS NULL ALTER TABLE dbo.admin_users ADD contact nvarchar(80) NULL;",
    "IF COL_LENGTH(N'dbo.admin_users',N'avatar_data') IS NULL ALTER TABLE dbo.admin_users ADD avatar_data varbinary(max) NULL;",
    "IF COL_LENGTH(N'dbo.admin_users',N'avatar_content_type') IS NULL ALTER TABLE dbo.admin_users ADD avatar_content_type varchar(100) NULL;",
    "IF COL_LENGTH(N'dbo.admin_users',N'must_change_password') IS NULL ALTER TABLE dbo.admin_users ADD must_change_password bit NOT NULL CONSTRAINT DF_admin_users_must_change DEFAULT(0) WITH VALUES;",
    "IF EXISTS(SELECT 1 FROM sys.check_constraints WHERE parent_object_id=OBJECT_ID(N'dbo.admin_users') AND name=N'CK_admin_users_role') ALTER TABLE dbo.admin_users DROP CONSTRAINT CK_admin_users_role;",
    "ALTER TABLE dbo.admin_users WITH CHECK ADD CONSTRAINT CK_admin_users_role CHECK(role IN ('ADMINISTRATOR','VIEWER','GROUP_ADMIN'));",
    "",
    "IF OBJECT_ID(N'dbo.admin_user_groups',N'U') IS NULL",
    "BEGIN",
    "  CREATE TABLE dbo.admin_user_groups(",
    "    user_id char(36) NOT NULL,group_name nvarchar(150) NOT NULL,region nvarchar(120) NULL,",
    "    assigned_at datetime2(3) NOT NULL CONSTRAINT DF_admin_user_groups_assigned DEFAULT(SYSUTCDATETIME()),",
    "    CONSTRAINT PK_admin_user_groups PRIMARY KEY(user_id,group_name),",
    "    CONSTRAINT FK_admin_user_groups_user FOREIGN KEY(user_id) REFERENCES dbo.admin_users(id) ON DELETE CASCADE",
    "  );",
    "END;",
    "IF NOT EXISTS(SELECT 1 FROM sys.indexes WHERE object_id=OBJECT_ID(N'dbo.admin_user_groups') AND name=N'IX_admin_user_groups_group') CREATE INDEX IX_admin_user_groups_group ON dbo.admin_user_groups(group_name,user_id);",
    "",
    "BEGIN TRANSACTION;",
    "CREATE TABLE #ImportedUsers(id char(36),email nvarchar(320),display_name nvarchar(160),region nvarchar(120),contact nvarchar(80),password_salt binary(16),password_hash binary(32));",
    "INSERT INTO #ImportedUsers(id,email,display_name,region,contact,password_salt,password_hash) VALUES",
]

admin_values = []
for admin in sorted(admins.values(), key=lambda item: item["email"]):
    admin_values.append(
        f"('{admin['id']}',{sql_text(admin['email'])},{sql_text(admin['name'])},{sql_text(admin['region'])},{sql_text(admin['contact'])},0x{admin['salt'].hex()},0x{admin['hash'].hex()})"
    )
sql.append(",\n".join(admin_values) + ";")
sql.extend([
    "",
    "MERGE dbo.admin_users WITH(HOLDLOCK) AS target",
    "USING #ImportedUsers AS source ON target.email=source.email",
    "WHEN MATCHED THEN UPDATE SET display_name=source.display_name,role='GROUP_ADMIN',region=source.region,contact=source.contact,is_active=1,updated_at=SYSUTCDATETIME()",
    "WHEN NOT MATCHED THEN INSERT(id,email,display_name,role,password_salt,password_hash,password_iterations,is_active,region,contact,must_change_password)",
    "VALUES(source.id,source.email,source.display_name,'GROUP_ADMIN',source.password_salt,source.password_hash,210000,1,source.region,source.contact,1);",
    "",
    "CREATE TABLE #Assignments(email nvarchar(320),group_name nvarchar(150),region nvarchar(120));",
    "INSERT INTO #Assignments(email,group_name,region) VALUES",
])

assignment_values = []
for admin in sorted(admins.values(), key=lambda item: item["email"]):
    for group in sorted(admin["groups"]):
        assignment_values.append(f"({sql_text(admin['email'])},{sql_text(group)},{sql_text(admin['region'])})")
if assignment_values:
    sql.append(",\n".join(assignment_values) + ";")
else:
    sql[-1] = "-- No se encontraron asignaciones." 
sql.extend([
    "",
    "DELETE ug FROM dbo.admin_user_groups ug INNER JOIN dbo.admin_users u ON u.id=ug.user_id INNER JOIN #ImportedUsers i ON i.email=u.email;",
    "INSERT INTO dbo.admin_user_groups(user_id,group_name,region)",
    "SELECT DISTINCT u.id,a.group_name,a.region FROM #Assignments a INNER JOIN dbo.admin_users u ON u.email=a.email INNER JOIN (SELECT DISTINCT grupo FROM dbo.agencies) g ON g.grupo=a.group_name;",
    "DROP TABLE #Assignments;",
    "DROP TABLE #ImportedUsers;",
    "COMMIT TRANSACTION;",
    "",
    "SELECT COUNT(*) AS administradores_de_grupo FROM dbo.admin_users WHERE role='GROUP_ADMIN' AND is_active=1;",
    "SELECT COUNT(*) AS asignaciones_activas FROM dbo.admin_user_groups;",
    "SELECT a.grupo FROM (SELECT DISTINCT grupo FROM dbo.agencies) a LEFT JOIN dbo.admin_user_groups ug ON ug.group_name=a.grupo WHERE ug.user_id IS NULL ORDER BY a.grupo;",
])
OUTPUT_SQL.write_text("\n".join(sql) + "\n", encoding="utf-8")

with CREDENTIALS_CSV.open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.writer(handle)
    writer.writerow(["Administrador", "Usuario", "Contraseña temporal", "Contacto", "Región", "Grupos asignados", "Agencias visibles", "Acción obligatoria"])
    for admin in sorted(admins.values(), key=lambda item: normalized(item["name"])):
        writer.writerow([admin["name"], admin["email"], admin["password"], admin["contact"], admin["region"], " | ".join(sorted(admin["groups"])), sum(group_counts[g] for g in admin["groups"]), "Cambiar contraseña en el primer ingreso"])

with ASSIGNMENTS_CSV.open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.writer(handle)
    writer.writerow(["Grupo SQL Server", "Agencias", "Administrador", "Usuario", "Región", "Estado"])
    for group in actual_groups:
        owner = group_owner.get(group)
        if owner:
            admin = admins[owner]
            writer.writerow([group, group_counts[group], admin["name"], admin["email"], admin["region"], "Asignado"])
        else:
            writer.writerow([group, group_counts[group], "Administrador principal", "admin@local", "Cobertura global", "Solo principal - revisar asignación"])

assigned = len(group_owner)
print(f"Generados {len(admins)} usuarios de grupo; {assigned}/{len(actual_groups)} grupos asignados y {len(actual_groups)-assigned} reservados al administrador principal.")
