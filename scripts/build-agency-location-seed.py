import json
import sys
from pathlib import Path

import openpyxl


def text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def code(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text(value)


def coordinate(value):
    if value is None:
        return None
    normalized = str(value).strip().rstrip(",").replace(",", ".")
    return float(normalized)


source = Path(sys.argv[1]).resolve()
destination = Path(sys.argv[2]).resolve()
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook.active

header_row = None
columns = {}
for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
    normalized = {
        str(value).strip().upper(): index
        for index, value in enumerate(row)
        if value is not None
    }
    if {"CODIGO", "DIRECCION", "LATITUDE", "LONGITUDE"}.issubset(normalized):
        header_row = row_number
        columns = normalized
        break

if header_row is None:
    raise SystemExit("No se encontraron los encabezados de agencias y coordenadas.")

records = []
seen = set()
for row_number, row in enumerate(
    sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
):
    agency_code = code(row[columns["CODIGO"]])
    if not agency_code:
        continue
    if agency_code in seen:
        raise SystemExit(f"Código duplicado en la fila {row_number}: {agency_code}")
    latitude = coordinate(row[columns["LATITUDE"]])
    longitude = coordinate(row[columns["LONGITUDE"]])
    if latitude is None or longitude is None or not (-90 <= latitude <= 90) or not (-180 <= longitude <= 180) or (latitude == 0 and longitude == 0):
        raise SystemExit(f"Coordenadas inválidas en la fila {row_number}: {agency_code}")
    seen.add(agency_code)
    records.append(
        {
            "codigo": agency_code,
            "direccion": text(row[columns["DIRECCION"]]),
            "sector": text(row[columns["SECTOR"]]) if "SECTOR" in columns else None,
            "municipio": text(row[columns["MUNICIPIO"]]) if "MUNICIPIO" in columns else None,
            "provincia": text(row[columns["PROVINCIA"]]) if "PROVINCIA" in columns else None,
            "latitude": latitude,
            "longitude": longitude,
        }
    )

if not records:
    raise SystemExit("El archivo no contiene agencias para sincronizar.")

destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(records, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(json.dumps({"source": source.name, "records": len(records), "destination": str(destination)}, ensure_ascii=False))
